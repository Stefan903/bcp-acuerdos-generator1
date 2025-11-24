from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
import psycopg2
from psycopg2 import sql
import os
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io
import logging

app = Flask(__name__)
CORS(app)  # Permitir peticiones desde GitHub Pages

logging.basicConfig(level=logging.INFO)

# Configuración de la base de datos
DB_CONFIG = {
    "host": os.environ.get("DB_HOST", "172.16.80.225"),
    "port": os.environ.get("DB_PORT", "5432"),
    "database": os.environ.get("DB_NAME", "ASIGNACION"),
    "user": os.environ.get("DB_USER", "postgres"),
    "password": os.environ.get("DB_PASSWORD", "912782014")
}

MESES = [
    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
    "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"
]

def get_db_connection():
    """Obtiene una conexión a la base de datos"""
    try:
        conn = psycopg2.connect(**DB_CONFIG)
        return conn
    except Exception as e:
        logging.error(f"Error al conectar a la base de datos: {str(e)}")
        raise

def determinar_tipo_documento(idc):
    """Determina el tipo de documento basado en el IDC"""
    idc_str = str(idc)
    if idc_str.startswith('1'):
        return "DNI"
    elif idc_str.startswith('3'):
        return "CE"
    elif idc_str.startswith('6'):
        return "RUC"
    else:
        return "DNI"

def formatear_numero_documento(idc, tipo_doc):
    """Formatea el número de documento según el tipo"""
    idc_str = str(idc)
    if tipo_doc == "DNI" and idc_str.startswith('1'):
        return idc_str[1:]
    elif tipo_doc == "CE" and idc_str.startswith('3'):
        return "0" + idc_str[1:]
    elif tipo_doc == "RUC" and idc_str.startswith('6'):
        return ""
    else:
        return idc_str

@app.route('/api/health', methods=['GET'])
def health_check():
    """Endpoint para verificar que el API está funcionando"""
    return jsonify({"status": "ok", "message": "API funcionando correctamente"})

@app.route('/api/buscar-cliente/<idc>', methods=['GET'])
def buscar_cliente(idc):
    """Busca un cliente y sus productos por IDC"""
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        
        query = """
        SELECT idc, nombre_titular, direccion, productorbp, numcuenta, 
               deudatotal, moneda, ubicacion
        FROM clientes_deudas 
        WHERE idc = %s
        ORDER BY productorbp
        """
        
        cursor.execute(query, (idc,))
        resultados = cursor.fetchall()
        
        if not resultados:
            return jsonify({"error": "Cliente no encontrado"}), 404
        
        primer_registro = resultados[0]
        tipo_doc = determinar_tipo_documento(idc)
        num_doc = formatear_numero_documento(idc, tipo_doc)
        
        cliente_info = {
            'idc': primer_registro[0],
            'nombre': primer_registro[1],
            'direccion': primer_registro[2],
            'distrito': primer_registro[7],
            'tipoDoc': tipo_doc,
            'numDoc': num_doc
        }
        
        productos = []
        for registro in resultados:
            moneda = "PEN" if registro[6] == "SOL" else "USD"
            producto = {
                'producto': registro[3],
                'cuenta': registro[4],
                'deuda': str(float(registro[5])),
                'cancelacion': str(float(registro[5])),
                'cuotas': '0',
                'moneda': moneda,
                'fechasPago': [datetime.now().strftime('%Y-%m-%d')],
                'montosPago': [str(float(registro[5]))]
            }
            productos.append(producto)
        
        cursor.close()
        conn.close()
        
        return jsonify({
            'cliente': cliente_info,
            'productos': productos
        })
        
    except Exception as e:
        logging.error(f"Error al buscar cliente: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/generar-documento', methods=['POST'])
def generar_documento():
    """Genera un documento Word con los datos del acuerdo"""
    try:
        data = request.json
        cliente = data.get('cliente')
        productos = data.get('productos')
        
        if not cliente or not productos:
            return jsonify({"error": "Datos incompletos"}), 400
        
        # Crear documento
        doc = Document()
        
        # Configurar márgenes
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)
        
        # Título
        titulo = doc.add_paragraph()
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = titulo.add_run("ACUERDO DE CANCELACIÓN")
        run.bold = True
        run.font.size = Pt(16)
        
        doc.add_paragraph()
        
        # Fecha
        fecha_actual = datetime.now()
        fecha_formateada = f"LIMA, {fecha_actual.day} DE {MESES[fecha_actual.month-1]} DEL {fecha_actual.year}"
        fecha_p = doc.add_paragraph(fecha_formateada)
        fecha_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()
        
        # Datos del cliente
        doc.add_paragraph(f"Señor(a): {cliente['nombre']}")
        doc.add_paragraph(f"Dirección: {cliente['direccion']}")
        doc.add_paragraph(f"Distrito: {cliente['distrito']}")
        doc.add_paragraph(f"Documento: {cliente['tipoDoc']}: {cliente['numDoc']}")
        
        doc.add_paragraph()
        
        # Tabla de productos
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Light Grid Accent 1'
        
        # Encabezados
        headers = table.rows[0].cells
        headers[0].text = 'Producto'
        headers[1].text = 'Cuenta'
        headers[2].text = 'Deuda'
        headers[3].text = 'Cancelación'
        headers[4].text = 'Cuotas'
        
        for header_cell in headers:
            for paragraph in header_cell.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
        
        # Agregar productos
        for prod in productos:
            row = table.add_row().cells
            simbolo = "S/" if prod['moneda'] == "PEN" else "USD"
            row[0].text = prod['producto']
            row[1].text = prod['cuenta']
            row[2].text = f"{simbolo} {float(prod['deuda']):,.2f}"
            row[3].text = f"{simbolo} {float(prod['cancelacion']):,.2f}"
            row[4].text = str(prod['cuotas'])
        
        doc.add_paragraph()
        
        # Forma de pago
        doc.add_paragraph(f"El día {cliente.get('fechaPago', datetime.now().strftime('%d/%m/%Y'))} por el importe de:")
        
        for prod in productos:
            simbolo = "S/" if prod['moneda'] == "PEN" else "USD"
            for fecha, monto in zip(prod['fechasPago'], prod['montosPago']):
                fecha_obj = datetime.strptime(fecha, '%Y-%m-%d')
                fecha_formateada = fecha_obj.strftime('%d/%m/%Y')
                doc.add_paragraph(
                    f"• {fecha_formateada} por el monto de {simbolo} {float(monto):,.2f} - ({prod['cuenta']})"
                )
        
        doc.add_paragraph()
        doc.add_paragraph("El cual será realizado a través de nuestra red o agencias a nivel nacional.")
        doc.add_paragraph("Este documento carece de valor sin el voucher de pago correspondiente.")
        doc.add_paragraph(
            "No olvides que, una vez cancelado el monto total de tu deuda, el plazo de condonación "
            "es de no menor a 15 días, luego de esto, podrás solicitar que se extienda tu constancia "
            "de no adeudo en un plazo no menor a 7 días."
        )
        
        # Guardar en memoria
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        
        filename = f"ACUERDO_CANCELACION_{cliente['nombre'].replace(' ', '_')}.docx"
        
        return send_file(
            file_stream,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            as_attachment=True,
            download_name=filename
        )
        
    except Exception as e:
        logging.error(f"Error al generar documento: {str(e)}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/test-db', methods=['GET'])
def test_db():
    """Prueba la conexión a la base de datos"""
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM clientes_deudas")
        count = cursor.fetchone()[0]
        cursor.close()
        conn.close()
        return jsonify({
            "status": "ok",
            "message": "Conexión exitosa",
            "registros": count
        })
    except Exception as e:
        return jsonify({
            "status": "error",
            "message": str(e)
        }), 500

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=True)
